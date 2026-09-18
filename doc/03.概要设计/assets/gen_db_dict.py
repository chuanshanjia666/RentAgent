#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按 shengyadb.xls 的体例生成 RentAgent 数据库表结构字典（一表一 sheet）。

体例（与参考文件逐格对应）：
  第 1 行  A1 = 表名
  第 2 行  B=表描述  C:D 合并=中文表描述  E:G 合并=做成日期  H:N 合并=日期
  第 3 行  B=表名称  C:D 合并=表名      E:G 合并=备注      H:N 合并=备注内容
  第 4 行  B..P 表头：类型|大小|名称|非空|默认值|补充|描述|扩展|左|类型|右|List显示|特殊取值描述|追加日期|追加原因
  第 5 行起 字段行；空 2 行后 A 列写 END

数据来源：docker/mysql-init/01_schema.sql（字段、类型、非空、默认值、外键、索引）
         与《数据库设计简介》4.3 数据字典（中文说明）。
"""
import pathlib
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

DO_DATE = "2026-09-16"
OUT = pathlib.Path(__file__).parent / "RentAgent-数据库表.xlsx"

HEADERS = ["类型", "大小", "名称", "非空", "默认值", "补充", "描述", "扩展",
           "左", "类型", "右", "List显示", "特殊取值描述", "追加日期", "追加原因"]
COL_WIDTHS = {"A": 3, "B": 12, "C": 8, "D": 26, "E": 6, "F": 16, "G": 34, "H": 30,
              "I": 8, "J": 6, "K": 10, "L": 6, "M": 9, "N": 46, "O": 11, "P": 12}


def F(ttype, size, name, notnull="", default="", note="", desc="",
      control="", list_show="", enum=""):
    """一个字段行。"""
    return dict(ttype=ttype, size=size, name=name, notnull=notnull, default=default,
                note=note, desc=desc, control=control, list_show=list_show, enum=enum)


# ── 18 张表 ────────────────────────────────────────────────────────────────
TABLES = []

TABLES.append(dict(
    name="sys_user", desc="用户表",
    remark="三类角色共用一张用户表（租客/房东/管理员）；username、phone 各建唯一索引；password 为 BCrypt 密文",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("VARCHAR", 50, "username", "Y", "", "UNIQUE uk_username", "账号", "input", "Y"),
        F("VARCHAR", 100, "password", "Y", "", "", "BCrypt 密文", "password"),
        F("VARCHAR", 20, "phone", "", "", "UNIQUE uk_phone", "手机号", "input", "Y"),
        F("VARCHAR", 100, "email", "", "", "", "邮箱", "input"),
        F("VARCHAR", 50, "nickname", "", "", "", "昵称", "input", "Y"),
        F("VARCHAR", 255, "avatar_url", "", "", "", "头像地址", "upload"),
        F("TINYINT", 1, "role", "Y", 1, "", "角色（简化 RBAC）", "select", "Y",
          "1：租客，2：房东，3：管理员"),
        F("TINYINT", 1, "status", "Y", 1, "", "账号状态（禁用即时生效）", "switch", "Y",
          "0：禁用，1：正常"),
        F("TINYINT", 1, "deleted", "Y", 0, "", "逻辑删除", "switch", "", "0：未删除，1：已删除"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "创建时间", "datetime", "Y"),
        F("DATETIME", "", "updated_at", "Y", "CURRENT_TIMESTAMP", "ON UPDATE CURRENT_TIMESTAMP",
          "更新时间", "datetime"),
    ]))

TABLES.append(dict(
    name="realname_auth", desc="实名认证表",
    remark="FR-03 实名认证；身份证号 AES 加密存储、另存 SHA-256 哈希用于同号比对；一人可多次提交，业务层取最新记录生效",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "user_id", "Y", "", "FK fk_realname_user → sys_user.id", "认证人", "select", "Y"),
        F("VARCHAR", 50, "real_name", "Y", "", "", "真实姓名", "input", "Y"),
        F("VARCHAR", 255, "id_card_no_enc", "Y", "", "AES 加密", "身份证号密文", "password"),
        F("CHAR", 64, "id_card_hash", "Y", "", "SHA-256", "身份证号哈希（比对用）"),
        F("TINYINT", 1, "status", "Y", 0, "", "认证状态", "select", "Y",
          "0：待审核，1：已通过，2：已驳回"),
        F("VARCHAR", 200, "reject_reason", "", "", "", "驳回理由", "textarea"),
        F("BIGINT", 20, "audit_by", "", "", "逻辑关联 sys_user.id", "审核人"),
        F("DATETIME", "", "audit_time", "", "", "", "审核时间", "datetime"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "提交时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="notification", desc="站内通知表",
    remark="FR-21 消息通知；审核结果、预约状态、签约、账单、举报处理等关键事件统一出口；索引 idx_user_read 支撑未读列表",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "user_id", "Y", "", "FK fk_notif_user → sys_user.id；KEY idx_user_read",
          "接收人", "select", "Y"),
        F("TINYINT", 1, "type", "Y", "", "", "通知类型", "select", "Y",
          "1：预约，2：审核，3：签约，4：账单，5：举报，9：系统"),
        F("VARCHAR", 100, "title", "Y", "", "", "通知标题", "input", "Y"),
        F("VARCHAR", 500, "content", "Y", "", "", "通知内容", "input", "Y"),
        F("VARCHAR", 30, "ref_type", "", "", "", "关联业务类型（house/appointment/contract…）", "input"),
        F("BIGINT", 20, "ref_id", "", "", "", "关联业务主键", "input"),
        F("TINYINT", 1, "is_read", "Y", 0, "", "已读标记", "switch", "Y", "0：未读，1：已读"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "触发时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="house", desc="房源表",
    remark="核心业务表；status 为审核状态机（FR-05/06/07）；索引：idx_status_district_rent、idx_status_layout、"
           "idx_landlord、idx_community、idx_lnglat（FR-09/10 检索域）",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "landlord_id", "Y", "", "FK fk_house_landlord → sys_user.id；KEY idx_landlord",
          "发布房东", "select", "Y"),
        F("VARCHAR", 100, "title", "Y", "", "", "标题（关键词检索）", "input", "Y"),
        F("VARCHAR", 100, "community", "Y", "", "KEY idx_community", "小区（定级分组维度）", "input", "Y"),
        F("VARCHAR", 50, "city", "Y", "", "", "城市", "select"),
        F("VARCHAR", 50, "district", "Y", "", "KEY idx_status_district_rent", "行政区", "select", "Y"),
        F("VARCHAR", 200, "address", "Y", "", "", "详细地址", "input"),
        F("VARCHAR", 20, "layout", "Y", "", "KEY idx_status_layout", "户型，如 2室1厅", "input", "Y"),
        F("DECIMAL", "6,2", "area", "Y", "", "", "面积（㎡）", "number", "Y"),
        F("VARCHAR", 10, "orientation", "", "", "", "朝向", "select", "Y"),
        F("VARCHAR", 20, "floor_desc", "", "", "", "楼层描述", "input"),
        F("DECIMAL", "10,2", "rent", "Y", "", "KEY idx_status_district_rent", "月租金", "number", "Y"),
        F("VARCHAR", 20, "deposit_type", "Y", "", "", "押付方式，如 押一付三", "select", "Y"),
        F("JSON", "", "facilities", "", "", "JSON 数组", "设施清单", "checkbox"),
        F("VARCHAR", 2000, "description", "", "", "", "房源描述（AI 辅助填充对象）", "textarea"),
        F("VARCHAR", 255, "cover_url", "", "", "", "封面图（列表卡片）", "upload"),
        F("DECIMAL", "10,7", "lng", "Y", "", "KEY idx_lnglat", "经度（地图找房）", "number"),
        F("DECIMAL", "10,7", "lat", "Y", "", "KEY idx_lnglat", "纬度（地图找房）", "number"),
        F("INT", 10, "view_count", "Y", 0, "UNSIGNED", "浏览数（推荐热度因子）", "number", "", ""),
        F("DECIMAL", "3,1", "avg_score", "", "", "", "平均评分（评价聚合冗余）", "", "Y"),
        F("TINYINT", 1, "status", "Y", 0, "", "房源状态（审核状态机）", "select", "Y",
          "0：待审核，1：已通过，2：已驳回，3：已上架，4：已下架，5：已出租"),
        F("VARCHAR", 200, "reject_reason", "", "", "", "驳回理由", "textarea"),
        F("TINYINT", 1, "deleted", "Y", 0, "", "逻辑删除", "switch", "", "0：未删除，1：已删除"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "创建时间", "datetime", "Y"),
        F("DATETIME", "", "updated_at", "Y", "CURRENT_TIMESTAMP", "ON UPDATE CURRENT_TIMESTAMP",
          "更新时间", "datetime"),
    ]))

TABLES.append(dict(
    name="house_image", desc="房源图片表",
    remark="房源与图片为整体—部分联系，随房源 ON DELETE CASCADE 级联物理删除",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "house_id", "Y", "", "FK fk_image_house → house.id ON DELETE CASCADE；KEY idx_house",
          "所属房源", "select", "Y"),
        F("VARCHAR", 255, "url", "Y", "", "", "图片地址", "upload", "Y"),
        F("TINYINT", 1, "sort", "Y", 0, "", "排序", "number"),
        F("TINYINT", 1, "is_cover", "Y", 0, "", "是否封面", "switch", "Y", "0：否，1：是"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "上传时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="favorite", desc="房源收藏表",
    remark="FR-25 收藏管理；用户与房源为 m:n 联系，由本表单独落地；UNIQUE(user_id, house_id) 防重复收藏",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "user_id", "Y", "", "FK fk_fav_user → sys_user.id；UNIQUE uk_user_house",
          "收藏人", "select", "Y"),
        F("BIGINT", 20, "house_id", "Y", "", "FK fk_fav_house → house.id；UNIQUE uk_user_house",
          "被收藏房源", "select", "Y"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "收藏时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="viewing_appointment", desc="看房预约表",
    remark="FR-17 预约看房；UNIQUE(active_slot) 防同时段重复预约——仅待确认/已确认持有占位键，"
           "终态置 NULL 释放时段；landlord_id 为查询冗余列（源于房源）",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "house_id", "Y", "", "FK fk_appt_house → house.id",
          "预约房源", "select", "Y"),
        F("BIGINT", 20, "tenant_id", "Y", "", "FK fk_appt_tenant → sys_user.id；KEY idx_tenant",
          "发起租客", "select", "Y"),
        F("BIGINT", 20, "landlord_id", "Y", "", "FK fk_appt_landlord → sys_user.id；KEY idx_landlord",
          "房东（查询冗余列）", "select"),
        F("DATETIME", "", "appointment_time", "Y", "", "", "预约时段（起点）", "datetime", "Y"),
        F("TINYINT", 1, "status", "Y", 0, "", "预约状态", "select", "Y",
          "0：待确认，1：已确认，2：已拒绝，3：已完成，4：已取消"),
        F("VARCHAR", 64, "active_slot", "", "", "UNIQUE uk_active_slot", "时段占位键", "input", "",
          "值形如 house_id:appointment_time，仅待确认/已确认非空，终态置 NULL 释放时段"),
        F("VARCHAR", 200, "reject_reason", "", "", "", "拒绝理由", "textarea"),
        F("VARCHAR", 200, "remark", "", "", "", "租客留言", "input"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "创建时间", "datetime", "Y"),
        F("DATETIME", "", "updated_at", "Y", "CURRENT_TIMESTAMP", "ON UPDATE CURRENT_TIMESTAMP",
          "更新时间", "datetime"),
    ]))

TABLES.append(dict(
    name="contract", desc="电子合同表",
    remark="FR-18 在线签约；clauses 为模板渲染后的条款集合（JSON，读多写少整体存取）；risk_flags 供 AI 解读标红；"
           "索引 idx_tenant_status、idx_landlord_status 支撑双方合同列表",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "house_id", "Y", "", "FK fk_contract_house → house.id", "签约房源", "select", "Y"),
        F("BIGINT", 20, "tenant_id", "Y", "", "FK fk_contract_tenant → sys_user.id；KEY idx_tenant_status",
          "租客", "select", "Y"),
        F("BIGINT", 20, "landlord_id", "Y", "", "FK fk_contract_landlord → sys_user.id；KEY idx_landlord_status",
          "房东", "select"),
        F("JSON", "", "clauses", "Y", "", "JSON 数组", "合同条款集合（模板渲染结果）", "textarea"),
        F("JSON", "", "risk_flags", "", "", "JSON 数组", "AI 解读风险条款标注"),
        F("DATE", "", "start_date", "Y", "", "", "租期起", "date", "Y"),
        F("DATE", "", "end_date", "Y", "", "", "租期止", "date", "Y"),
        F("DECIMAL", "10,2", "monthly_rent", "Y", "", "", "月租金", "number", "Y"),
        F("DECIMAL", "10,2", "deposit", "Y", "", "", "押金", "number", "Y"),
        F("TINYINT", 1, "status", "Y", 0, "", "合同状态（双方签署状态机）", "select", "Y",
          "0：待租客确认，1：待房东确认，2：已生效，3：已退租，4：已作废"),
        F("DATETIME", "", "signed_tenant_at", "", "", "", "租客签署时间", "datetime"),
        F("DATETIME", "", "signed_landlord_at", "", "", "", "房东签署时间", "datetime"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "创建时间", "datetime", "Y"),
        F("DATETIME", "", "updated_at", "Y", "CURRENT_TIMESTAMP", "ON UPDATE CURRENT_TIMESTAMP",
          "更新时间", "datetime"),
    ]))

TABLES.append(dict(
    name="lease_order", desc="租赁订单表",
    remark="FR-19 订单管理；合同与订单为 1:1，UNIQUE(contract_id) 落地；合同生效时在同一事务内派生",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "contract_id", "Y", "", "FK fk_order_contract → contract.id；UNIQUE uk_contract（1:1）",
          "来源合同", "select", "Y"),
        F("BIGINT", 20, "house_id", "Y", "", "FK fk_order_house → house.id", "房源", "select"),
        F("BIGINT", 20, "tenant_id", "Y", "", "FK fk_order_tenant → sys_user.id；KEY idx_tenant", "租客", "select", "Y"),
        F("BIGINT", 20, "landlord_id", "Y", "", "FK fk_order_landlord → sys_user.id；KEY idx_landlord",
          "房东", "select"),
        F("DATE", "", "start_date", "Y", "", "", "租期起", "date", "Y"),
        F("DATE", "", "end_date", "Y", "", "", "租期止", "date", "Y"),
        F("DECIMAL", "10,2", "monthly_rent", "Y", "", "", "月租金", "number", "Y"),
        F("DECIMAL", "10,2", "deposit", "Y", "", "", "押金", "number", "Y"),
        F("TINYINT", 1, "status", "Y", 0, "", "订单状态", "select", "Y",
          "0：在租，1：已退租，2：已到期"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "创建时间", "datetime", "Y"),
        F("DATETIME", "", "updated_at", "Y", "CURRENT_TIMESTAMP", "ON UPDATE CURRENT_TIMESTAMP",
          "更新时间", "datetime"),
    ]))

TABLES.append(dict(
    name="rent_bill", desc="租金账单表",
    remark="FR-19 账单计划；UNIQUE(lease_order_id, period_no) 保证按月生成计划幂等；仅记录，不对接真实支付渠道",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "lease_order_id", "Y", "", "FK fk_bill_order → lease_order.id；UNIQUE uk_order_period",
          "所属订单", "select", "Y"),
        F("INT", 10, "period_no", "Y", "", "UNIQUE uk_order_period", "第几期", "number", "Y"),
        F("DATE", "", "due_date", "Y", "", "", "应付日", "date", "Y"),
        F("DECIMAL", "10,2", "amount", "Y", "", "", "应付金额", "number", "Y"),
        F("TINYINT", 1, "status", "Y", 0, "", "账单状态", "select", "Y",
          "0：待支付，1：已支付，2：已逾期"),
        F("DATETIME", "", "paid_at", "", "", "", "支付标记时间", "datetime", "Y"),
        F("VARCHAR", 200, "remark", "", "", "", "备注", "input"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "生成时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="review", desc="评价表",
    remark="FR-20 评价体系；UNIQUE(lease_order_id) 实现一单一评（仅完成合同可评）；"
           "CHECK 约束限定评分 1~5；索引 idx_house_status 支撑房源评价列表",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "lease_order_id", "Y", "", "FK fk_review_order → lease_order.id；UNIQUE uk_order（一单一评）",
          "评价来源订单", "select", "Y"),
        F("BIGINT", 20, "house_id", "Y", "", "FK fk_review_house → house.id；KEY idx_house_status",
          "被评房源", "select", "Y"),
        F("BIGINT", 20, "landlord_id", "Y", "", "FK fk_review_landlord → sys_user.id", "被评房东", "select"),
        F("BIGINT", 20, "tenant_id", "Y", "", "FK fk_review_tenant → sys_user.id", "评价人", "select"),
        F("TINYINT", 1, "house_score", "Y", "", "CHECK ck_house_score（1~5）", "房源评分", "number", "Y",
          "1 ~ 5 的整数"),
        F("TINYINT", 1, "landlord_score", "Y", "", "CHECK ck_landlord_score（1~5）", "房东评分", "number", "Y",
          "1 ~ 5 的整数"),
        F("VARCHAR", 500, "content", "", "", "", "文字评价", "textarea", "Y"),
        F("TINYINT", 1, "status", "Y", 0, "", "评价状态", "select", "Y", "0：正常，1：已隐藏"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "评价时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="report", desc="举报表",
    remark="FR-23 内容审核管理；target_type + target_id 为多态关联（不建物理外键，业务层校验）；"
           "索引 idx_status 支撑待处理列表",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "reporter_id", "Y", "", "FK fk_report_reporter → sys_user.id", "举报人", "select", "Y"),
        F("TINYINT", 1, "target_type", "Y", "", "多态关联", "举报对象类型", "select", "Y",
          "1：房源，2：评价，3：用户"),
        F("BIGINT", 20, "target_id", "Y", "", "多态关联（业务层校验）", "举报对象主键", "number"),
        F("VARCHAR", 200, "reason", "Y", "", "", "举报理由", "textarea", "Y"),
        F("TINYINT", 1, "status", "Y", 0, "KEY idx_status", "处理状态", "select", "Y",
          "0：待处理，1：已处理"),
        F("BIGINT", 20, "handle_by", "", "", "逻辑关联 sys_user.id", "处理人"),
        F("VARCHAR", 200, "handle_remark", "", "", "", "处理说明", "textarea"),
        F("DATETIME", "", "handled_at", "", "", "", "处理时间", "datetime"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "举报时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="audit_log", desc="审计日志表",
    remark="FR-07/22/23 操作留痕；只增不改（无 updated_at）；operator_id 与 target 为逻辑关联，"
           "索引 idx_operator、idx_target",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "operator_id", "Y", "", "逻辑关联 sys_user.id；KEY idx_operator", "操作人", "select", "Y"),
        F("VARCHAR", 50, "action", "Y", "", "", "动作标识", "input", "Y",
          "HOUSE_AUDIT / USER_BAN / REPORT_HANDLE / SIGN 等"),
        F("VARCHAR", 30, "target_type", "Y", "", "KEY idx_target（多态关联）", "操作对象类型", "select", "Y"),
        F("BIGINT", 20, "target_id", "Y", "", "KEY idx_target（多态关联）", "操作对象主键", "number", "Y"),
        F("JSON", "", "detail", "", "", "JSON 对象", "操作前后快照"),
        F("VARCHAR", 45, "ip", "", "", "兼容 IPv6", "操作 IP", "input"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "操作时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="ai_chat_session", desc="AI 会话表",
    remark="FR-12/13/14 会话载体；scene 决定角色设定与工具集路由；多轮记忆不单独建模——"
           "由 ai_chat_message 窗口读取 + context_summary 摘要压缩实现",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "user_id", "Y", "", "逻辑关联 sys_user.id；KEY idx_user（user_id, updated_at）",
          "会话归属", "select", "Y"),
        F("TINYINT", 1, "scene", "Y", "", "", "会话场景（成员路由依据）", "select", "Y",
          "1：找房助手，2：智能客服，3：合同解读"),
        F("VARCHAR", 100, "title", "", "", "", "会话标题（首条消息截取）", "input", "Y"),
        F("VARCHAR", 500, "context_summary", "", "", "", "长对话记忆摘要（窗口压缩）", "textarea"),
        F("TINYINT", 1, "is_transferred", "Y", 0, "", "客服未命中转人工标记", "switch", "Y",
          "0：未转人工，1：已转人工"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "创建时间", "datetime", "Y"),
        F("DATETIME", "", "updated_at", "Y", "CURRENT_TIMESTAMP", "ON UPDATE CURRENT_TIMESTAMP",
          "最后活跃时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="ai_chat_message", desc="AI 对话消息表",
    remark="NFR-05 对话留痕；role=3 时记录工具调用链（tool_name/tool_args/tool_result）实现 NFR-09 可追溯；"
           "latency_ms 用于 NFR-02 首字延迟度量",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "session_id", "Y", "", "FK fk_msg_session → ai_chat_session.id；KEY idx_session（session_id, id）",
          "所属会话", "select"),
        F("TINYINT", 1, "role", "Y", "", "", "消息角色", "select", "Y",
          "1：user，2：assistant，3：tool"),
        F("MEDIUMTEXT", "", "content", "", "", "", "消息正文", "textarea", "Y"),
        F("VARCHAR", 50, "tool_name", "", "", "", "工具名（role=3 时有值）", "input"),
        F("JSON", "", "tool_args", "", "", "JSON 对象", "工具入参"),
        F("JSON", "", "tool_result", "", "", "JSON 对象", "工具返回结果"),
        F("JSON", "", "citations", "", "", "JSON 数组", "RAG 来源引用（NFR-05）"),
        F("INT", 10, "token_count", "", "", "", "token 用量", "number"),
        F("INT", 10, "latency_ms", "", "", "", "响应延迟（毫秒）", "number", "Y"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "发生时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="ai_analysis", desc="AI 分析结果表",
    remark="FR-08/14/15/16 四类分析结果统一落本表；input_snapshot 保留输入快照作为复现依据；"
           "target 为多态逻辑关联，索引 idx_target、idx_user",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "user_id", "Y", "", "逻辑关联 sys_user.id；KEY idx_user", "触发人", "select", "Y"),
        F("TINYINT", 1, "type", "Y", "", "", "分析类型", "select", "Y",
          "1：定价建议，2：虚假房源检测，3：合同解读，4：房源信息识别"),
        F("VARCHAR", 30, "target_type", "Y", "", "KEY idx_target（多态关联）", "分析对象类型", "select", "Y"),
        F("BIGINT", 20, "target_id", "Y", "", "KEY idx_target（多态关联）", "分析对象主键", "number", "Y"),
        F("JSON", "", "input_snapshot", "", "", "JSON 对象", "输入快照（复现依据）"),
        F("JSON", "", "result", "Y", "", "JSON 对象", "结构化结果（区间/依据/样本量、风险分/疑点、条款解读）"),
        F("VARCHAR", 100, "model", "", "", "", "生成模型标识（协议:模型，换模型/换供应商可追溯）", "input", "Y"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "生成时间", "datetime", "Y"),
    ]))

TABLES.append(dict(
    name="kb_document", desc="知识库文档表",
    remark="FR-13 智能客服知识源；切片与向量重建以 content 为源头（NFR-10 不纳入备份关键路径）；"
           "索引 idx_category_status",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("VARCHAR", 100, "title", "Y", "", "", "文档标题", "input", "Y"),
        F("TINYINT", 1, "category", "Y", "", "KEY idx_category_status", "文档分类", "select", "Y",
          "1：平台规则，2：租赁政策 FAQ，3：合同模板说明"),
        F("VARCHAR", 255, "source_url", "", "", "", "原文出处", "input"),
        F("MEDIUMTEXT", "", "content", "Y", "", "", "正文（切片与向量重建源头）", "textarea"),
        F("TINYINT", 1, "status", "Y", 0, "", "入库状态", "select", "Y",
          "0：待入库，1：已生效，2：失效"),
        F("INT", 10, "chunk_count", "Y", 0, "", "切片数（冗余统计）", "number", "Y"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "创建时间", "datetime", "Y"),
        F("DATETIME", "", "updated_at", "Y", "CURRENT_TIMESTAMP", "ON UPDATE CURRENT_TIMESTAMP",
          "更新时间", "datetime"),
    ]))

TABLES.append(dict(
    name="kb_chunk", desc="知识库切片表",
    remark="RAG 检索与引用溯源的最小单位；UNIQUE(document_id, seq) 定位切片顺序；"
           "vector_ref 为 Redis Stack 向量索引引用位，随文档 ON DELETE CASCADE",
    fields=[
        F("BIGINT", 20, "id", "Y", "", "主键 自增 UNSIGNED", "主键"),
        F("BIGINT", 20, "document_id", "Y", "",
          "FK fk_chunk_document → kb_document.id ON DELETE CASCADE；UNIQUE uk_doc_seq",
          "所属文档", "select"),
        F("INT", 10, "seq", "Y", "", "UNIQUE uk_doc_seq", "切片序号", "number", "Y"),
        F("TEXT", "", "content", "Y", "", "", "切片正文（embedding 输入）", "textarea", "Y"),
        F("INT", 10, "token_count", "", "", "", "切片长度", "number"),
        F("VARCHAR", 64, "vector_ref", "", "", "向量库引用", "Redis Stack 向量引用 id（升级后填充）", "input"),
        F("DATETIME", "", "created_at", "Y", "CURRENT_TIMESTAMP", "", "入库时间", "datetime", "Y"),
    ]))


# ── 样式 ────────────────────────────────────────────────────────────────────
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
LABEL_FILL = PatternFill("solid", fgColor="F2F2F2")
NAME_FONT = Font(name="Consolas", size=10)
CN_FONT = Font(name="宋体", size=10)
BOLD = Font(name="宋体", size=10, bold=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_sheet(wb, t, idx):
    ws = wb.create_sheet(t["name"])
    for col, w in COL_WIDTHS.items():
        ws.column_dimensions[col].width = w

    # 第 1 行：表名
    ws["A1"] = t["name"]
    ws["A1"].font = Font(name="Consolas", size=12, bold=True)

    # 第 2 行：表描述 / 做成日期
    ws["B2"] = "表描述"; ws["B2"].font = BOLD; ws["B2"].fill = LABEL_FILL; ws["B2"].alignment = CENTER
    ws.merge_cells("C2:D2"); ws["C2"] = t["desc"]; ws["C2"].font = CN_FONT; ws["C2"].alignment = CENTER
    ws.merge_cells("E2:G2"); ws["E2"] = "做成日期"; ws["E2"].font = BOLD
    ws["E2"].fill = LABEL_FILL; ws["E2"].alignment = CENTER
    ws.merge_cells("H2:N2"); ws["H2"] = DO_DATE; ws["H2"].font = CN_FONT; ws["H2"].alignment = CENTER

    # 第 3 行：表名称 / 备注
    ws["B3"] = "表名称"; ws["B3"].font = BOLD; ws["B3"].fill = LABEL_FILL; ws["B3"].alignment = CENTER
    ws.merge_cells("C3:D3"); ws["C3"] = t["name"]; ws["C3"].font = NAME_FONT; ws["C3"].alignment = CENTER
    ws.merge_cells("E3:G3"); ws["E3"] = "备注"; ws["E3"].font = BOLD
    ws["E3"].fill = LABEL_FILL; ws["E3"].alignment = CENTER
    ws.merge_cells("H3:N3"); ws["H3"] = t["remark"]; ws["H3"].font = CN_FONT
    ws["H3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 42
    for rng in ("B2:G3",):
        pass

    # 第 4 行：表头
    for i, h in enumerate(HEADERS):
        c = ws.cell(row=4, column=2 + i, value=h)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[4].height = 20

    # 字段行
    r = 5
    for f in t["fields"]:
        vals = [f["ttype"], f["size"], f["name"], f["notnull"], f["default"],
                f["note"], f["desc"], "", "", f["control"], "", f["list_show"], f["enum"], "", ""]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=2 + i, value=(v if v != "" else None))
            c.border = BORDER
            if i == 2:                      # 名称列等宽字体，便于核对
                c.font = NAME_FONT
                c.alignment = LEFT
            elif i in (0, 1, 3, 4, 9, 11):
                c.font = CN_FONT
                c.alignment = CENTER
            else:
                c.font = CN_FONT
                c.alignment = LEFT
        r += 1

    # 结束行（与参考体例一致：末行字段后留两行空白再写 END）
    ws.cell(row=r + 2, column=1, value="END").font = BOLD
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False

    # 打印设置：横向、按宽适配一页、重复表头行，便于直接打印装订
    last = r + 2
    ws.print_area = f"A1:P{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "4:4"
    return ws


def build_index(wb):
    ws = wb.create_sheet("表清单", 0)
    headers = ["序号", "表名", "中文名", "业务域", "字段数", "主键", "外键", "唯一约束", "说明"]
    widths = [6, 26, 16, 10, 8, 8, 34, 34, 46]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i, value=h)
        c.font = BOLD; c.fill = HEAD_FILL; c.alignment = CENTER; c.border = BORDER

    domain = {"sys_user": "用户域", "realname_auth": "用户域", "notification": "用户域",
              "house": "房源域", "house_image": "房源域", "favorite": "房源域",
              "viewing_appointment": "交易域", "contract": "交易域", "lease_order": "交易域",
              "rent_bill": "交易域", "review": "交易域", "report": "交易域", "audit_log": "交易域",
              "ai_chat_session": "AI 域", "ai_chat_message": "AI 域", "ai_analysis": "AI 域",
              "kb_document": "AI 域", "kb_chunk": "AI 域"}

    for n, t in enumerate(TABLES, start=1):
        fks = [f["note"].split("→")[1].split(".")[0].strip()
               for f in t["fields"] if f["note"].startswith("FK") and "→" in f["note"]]
        uniq = [f["name"] for f in t["fields"] if "UNIQUE" in f["note"]]
        row = [n, t["name"], t["desc"], domain[t["name"]], len(t["fields"]),
               "id", "，".join(dict.fromkeys(fks)) or "—",
               "，".join(dict.fromkeys(uniq)) or "—", t["remark"]]
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=n + 1, column=i, value=v)
            c.border = BORDER
            c.font = NAME_FONT if i == 2 else CN_FONT
            c.alignment = CENTER if i in (1, 3, 4, 5, 6) else LEFT

    ws.cell(row=len(TABLES) + 3, column=1,
            value="说明：本工作簿共 18 张业务表，与《数据库设计简介》4.3 数据字典、"
                  "docker/mysql-init/01_schema.sql 三者一致；改字段请以 schema.sql 为准并同步更新本表。").font = BOLD
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.print_area = f"A1:I{len(TABLES) + 3}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"
    return ws


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_index(wb)
    for i, t in enumerate(TABLES):
        build_sheet(wb, t, i)
    wb.save(OUT)
    print("written:", OUT, "| sheets:", len(wb.sheetnames))
    print("tables:", len(TABLES), "| fields total:", sum(len(t["fields"]) for t in TABLES))


if __name__ == "__main__":
    main()
