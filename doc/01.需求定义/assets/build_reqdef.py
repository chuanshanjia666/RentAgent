#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 RentAgent《需求定义说明书》docx（交付物目录 01.需求定义）。

策略与《概要设计》一致：复用学校模板的样式表、页面设置、页眉页脚与封面/修改控制/目录区，
重建正文，再用 PDF 两遍校准目录页码与总页数。

版式构件（Builder / 目录重建 / 页码校准）从 03.概要设计 的生成脚本复用，避免两套实现漂移；
因此本脚本需与 `doc/03.概要设计/assets/` 同时存在。

运行：python3 build_reqdef.py            （需 libreoffice、python-docx、pymupdf、Pillow）
"""
import pathlib
import sys

BASE = pathlib.Path(__file__).resolve().parent
RENT = BASE.parents[1]                       # repo/doc
ASSET3 = RENT / "03.概要设计" / "assets"      # 版式构件来源
sys.path.insert(0, str(ASSET3))

import build_doc as BD                                        # noqa: E402
from build_doc import Builder, rebuild_toc, set_core_props, set_para_runs  # noqa: E402

import content_reqdef as C                                     # noqa: E402

BD.BASE = BASE                 # 让页码校准的临时 PDF 落在本目录
BD.FIG = BASE / "figures"      # 本目录图源
TPL = BASE / "tpl.docx"
OUT = BASE.parent / "需求定义说明书.docx"

USABLE = 16.99


# ── 主流程 ──────────────────────────────────────────────────────────────────
def build(out_path, toc_pages=None, total_pages=None):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document(str(TPL))
    set_core_props(doc, C.DOC_TITLE)
    body = doc.element.body

    # 1. 封面
    cover = ["密级：秘密",
             f"文档编号：{C.DOC_NO}",
             f"项目名称：{C.PROJECT}",
             f"项目编号：{C.PROJ_NO}",
             "", "", "", "",
             C.DOC_TITLE,
             f"版本：{C.DOC_VER}",
             C.DOC_DATE,
             "", "", "", "",
             C.ORG,
             "(版权所有，翻版必究)"]
    for i, txt in enumerate(cover):
        set_para_runs(doc.paragraphs[i]._p, txt)

    # 压缩封面空段落，使版权行与页数表并入首页（与概要设计同法）
    cover_paras = [doc.paragraphs[i]._p for i in range(18)]
    empties = []
    for el in cover_paras:
        txt = "".join(t.text or "" for t in el.iter(qn("w:t")))
        if txt.strip() == "":
            empties.append(el)
    from docx.text.paragraph import Paragraph
    for el in empties:
        pp = Paragraph(el, doc.paragraphs[0]._parent)
        pp.paragraph_format.space_before = Pt(0)
        pp.paragraph_format.space_after = Pt(0)
        pp.paragraph_format.line_spacing = Pt(7)
        for r in el.findall(qn("w:r")):
            rPr = r.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                r.insert(0, rPr)
            for tag in ("sz", "szCs"):
                e = rPr.find(qn("w:" + tag))
                if e is None:
                    e = OxmlElement("w:" + tag)
                    rPr.append(e)
                e.set(qn("w:val"), "12")
    for el in [empties[1], empties[2], empties[6], empties[7]]:
        el.getparent().remove(el)

    # 2. 文件修改控制表
    t = doc.tables[1]
    while len(t.rows) > 1:
        t._tbl.remove(t.rows[-1]._tr)
    for r in C.CHANGE_ROWS:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            set_para_runs(cells[i].paragraphs[0]._p, v, size=9)

    # 3. 页眉
    for sec in doc.sections:
        try:
            hdr = sec.header
        except Exception:
            continue
        for p in hdr.paragraphs:
            for r in p.runs:
                if "概要设计书" in r.text:
                    r.text = (f"RentAgent {C.DOC_TITLE}" + " " * 56 + "版本：")
                elif r.text.strip() == "0.1.0":
                    r.text = C.DOC_VER

    # 4. 清空模板正文（保留封面 / 修改控制 / 目录域与末尾 sectPr）
    children = list(body)
    for el in children[26:-1]:
        body.remove(el)

    # 5. 重建正文
    b = Builder(doc)
    section1(b)
    section2(b)
    section3(b)
    section4(b)
    section5(b)
    section6(b)
    section7(b)
    section8(b)
    section9(b)

    # 6. 目录
    rebuild_toc(doc, b.toc, toc_pages)

    # 7. 封面页数与签署
    if total_pages:
        t0 = doc.tables[0]
        vals = {"总页数": str(total_pages), "正文": str(max(total_pages - 4, 0)),
                "附录": "无", "生效日期": C.DOC_DATE.replace("-", ".")}
        for row in t0.rows:
            for i, cell in enumerate(row.cells):
                key = cell.text.strip()
                if key in vals and i + 1 < len(row.cells):
                    set_para_runs(row.cells[i + 1].paragraphs[0]._p, vals[key], size=8)
                if key == "编制" and i + 1 < len(row.cells):
                    set_para_runs(row.cells[i + 1].paragraphs[0]._p, C.AUTHOR)
                if key == "批准" and i + 1 < len(row.cells):
                    set_para_runs(row.cells[i + 1].paragraphs[0]._p, C.APPROVER)

    doc.save(str(out_path))
    return b.toc


# ── 第 1 章 ─────────────────────────────────────────────────────────────────
def section1(b):
    b.heading(1, "1", "文档概述", break_before=True)
    b.heading(2, "1.1", "文档目的和范围")
    b.body_many(C.SEC1_PURPOSE)
    b.heading(2, "1.2", "术语与缩略语")
    b.body("本文档使用的术语与缩略语如下表所示。")
    b.tcaption("表1-1 术语与缩略语")
    b.table(["序号", "术语/缩略语", "说明"], C.TERMS, [1.2, 3.6, 12.19],
            align_center_cols=(0,))
    b.heading(2, "1.3", "参考文档")
    b.body("本文档编写所依据或引用的文档如下表所示。")
    b.tcaption("表1-2 参考文档")
    b.table(["序号", "文档名", "作者", "时间", "版本"], C.REF_DOCS,
            [1.1, 6.9, 3.0, 2.5, 3.49], align_center_cols=(0, 3, 4))


# ── 第 2 章 ─────────────────────────────────────────────────────────────────
def section2(b):
    b.heading(1, "2", "项目背景与问题定义", break_before=True)
    b.heading(2, "2.1", "行业背景")
    b.body_many(C.SEC2_BG)

    b.heading(2, "2.2", "现有平台的主要不足")
    b.body_many(C.SEC2_GAP)
    b.tcaption("表2-1 现有租赁平台的主要不足")
    b.table(["序号", "痛点", "具体表现", "平台侧成因"], C.TABLE_2_1,
            [1.1, 2.6, 8.0, 5.29], align_center_cols=(0,))

    b.heading(2, "2.3", "问题定义")
    b.body_many(C.SEC2_DEFINE)
    b.figure("fig_2_1_pain", "图2-1  业务痛点与系统应对映射图", width_cm=16.2)

    b.heading(2, "2.4", "AI 智能体技术的适用性")
    b.body_many(C.SEC2_AI)


# ── 第 3 章 ─────────────────────────────────────────────────────────────────
def section3(b):
    b.heading(1, "3", "项目目标与范围", break_before=True)
    b.heading(2, "3.1", "总体目标")
    b.body_many(C.SEC3_GOAL)

    b.heading(2, "3.2", "可度量目标")
    b.body("项目的可度量目标如下表所示，验收时逐条核对。")
    b.tcaption("表3-1 可度量目标")
    b.table(["编号", "目标", "度量内容", "口径说明"], C.TABLE_3_1,
            [1.3, 3.3, 7.0, 5.39], align_center_cols=(0,))

    b.heading(2, "3.3", "系统范围（本期实现）")
    b.body_many(C.SEC3_SCOPE)
    b.tcaption("表3-2 本期范围内功能清单")
    b.table(["序号", "功能模块", "范围界定", "对应需求编号"], C.TABLE_3_2,
            [1.1, 2.7, 10.3, 2.89], align_center_cols=(0,))

    b.heading(2, "3.4", "范围外（本期不做）")
    b.body("为避免范围蔓延，下表内容明确列在本期范围之外，不在系统中实现。")
    b.tcaption("表3-3 本期范围外清单")
    b.table(["序号", "范围外内容", "本期处理方式", "排除理由"], C.TABLE_3_3,
            [1.1, 4.0, 6.4, 5.49], align_center_cols=(0,))

    b.heading(2, "3.5", "假设与约束")
    b.body("本文档成立所依赖的假设与项目必须遵守的约束如下表所示；假设不成立时需重新评估相应目标。")
    b.tcaption("表3-4 假设与约束")
    b.table(["类型", "编号", "内容", "影响"], C.TABLE_3_4,
            [1.4, 1.2, 10.4, 3.99], align_center_cols=(0, 1))


# ── 第 4 章 ─────────────────────────────────────────────────────────────────
def section4(b):
    b.heading(1, "4", "用户角色与业务场景", break_before=True)
    b.heading(2, "4.1", "用户角色定义")
    b.body_many(C.SEC4_ROLE)
    b.tcaption("表4-1 用户角色定义")
    b.table(["序号", "角色", "定位", "核心诉求", "权限与前置约定"], C.TABLE_4_1,
            [1.0, 2.0, 2.6, 6.4, 4.99], align_center_cols=(0,))

    b.heading(2, "4.2", "核心业务流程")
    b.body_many(C.SEC4_FLOW)
    b.figure("fig_4_1_flow", "图4-1  核心业务流程图", width_cm=13.5)

    b.heading(2, "4.3", "角色与用例")
    b.body_many(C.SEC4_UC)
    b.figure("fig_4_2_usecase", "图4-2  角色—用例对照图", width_cm=15.5)

    b.heading(2, "4.4", "典型业务场景")
    b.body_many(C.SEC4_SCENE)
    b.tcaption("表4-2 典型业务场景")
    b.table(["编号", "角色", "场景", "场景描述", "涉及需求"], C.TABLE_4_2,
            [1.0, 2.0, 2.8, 8.2, 2.99], size=8.5, align_center_cols=(0,))

    b.heading(2, "4.5", "关键业务规则")
    b.body_many(C.SEC4_RULES)
    for i, it in enumerate(C.SEC4_RULES_LIST, start=1):
        b.para(f"（{i}）{it}", style="正文缩进")


# ── 第 5 章 ─────────────────────────────────────────────────────────────────
def section5(b):
    b.heading(1, "5", "系统上下文与总体方案", break_before=True)
    b.heading(2, "5.1", "系统上下文")
    b.body_many(C.SEC5_CONTEXT)
    b.figure("fig_5_1_context", "图5-1  系统上下文图", width_cm=14.5)

    b.heading(2, "5.2", "系统总体架构")
    b.body_many(C.SEC5_ARCH)
    b.figure("fig_5_2_arch", "图5-2  系统总体架构图", width_cm=15.0)

    # 两张表各自独占一页，避免 9 行选型表被拦腰截断后留下只有两行的近空白页
    b.heading(2, "5.3", "技术方案选型", break_before=True)
    b.body("系统各层的技术选型与理由如下表所示。选型以“成熟、可复用、团队有基础”为原则，"
           "不引入需要长期学习成本的前沿组件。")
    b.tcaption("表5-1 技术方案选型")
    b.table(["层次", "选型", "选型理由"], C.TABLE_5_1, [2.2, 5.4, 9.39])

    # 5.3 的选型表占满前半页，5.4 的 AI 能力表整表另起一页，避免末行被挤成整页
    b.heading(2, "5.4", "AI 智能体能力定义", break_before=True)
    b.body_many(C.SEC5_AI_CAP)
    b.tcaption("表5-2 AI 智能体能力定义")
    b.table(["序号", "能力", "优先级", "能力说明与边界", "承载需求"], C.TABLE_5_2,
            [1.0, 3.0, 1.3, 9.2, 2.49], size=8.5, align_center_cols=(0, 2, 4))


# ── 第 6 章 ─────────────────────────────────────────────────────────────────
def section6(b):
    b.heading(1, "6", "可行性分析", break_before=True)
    b.heading(2, "6.1", "技术可行性")
    b.body_many(C.SEC6_TECH)

    b.heading(2, "6.2", "经济可行性")
    b.body_many(C.SEC6_ECON)
    b.tcaption("表6-1 资源与预算估算")
    b.table(["类别", "项目", "估算金额", "说明"], C.TABLE_6_1,
            [1.7, 5.2, 3.3, 6.79])

    b.heading(2, "6.3", "操作可行性")
    b.body_many(C.SEC6_OP)

    b.heading(2, "6.4", "法律与合规可行性")
    b.body_many(C.SEC6_LAW)

    b.heading(2, "6.5", "可行性结论")
    b.body("四个维度的分析结论汇总如下表所示，结论均为可行，项目具备启动条件。")
    b.tcaption("表6-2 可行性分析结论")
    b.table(["维度", "结论", "主要依据"], C.TABLE_6_2, [2.4, 1.5, 13.09],
            align_center_cols=(0, 1))


# ── 第 7 章 ─────────────────────────────────────────────────────────────────
def section7(b):
    b.heading(1, "7", "风险识别与应对", break_before=True)
    b.body_many(C.SEC7_INTRO)
    b.tcaption("表7-1 风险识别与应对措施")
    b.table(["编号", "风险", "等级", "应对措施", "责任角色"], C.TABLE_7_1,
            [1.1, 4.0, 1.2, 8.2, 2.49], size=8.5, align_center_cols=(0, 2))


# ── 第 8 章 ─────────────────────────────────────────────────────────────────
def section8(b):
    b.heading(1, "8", "开发计划与团队分工", break_before=True)
    b.heading(2, "8.1", "里程碑计划")
    b.body_many(C.SEC8_PLAN)
    b.tcaption("表8-1 阶段里程碑计划")
    b.table(["阶段", "阶段名称", "主要工作", "阶段成果物", "当前状态"], C.TABLE_8_1,
            [1.0, 2.0, 4.4, 5.4, 4.19], size=8.5, align_center_cols=(0,))

    b.heading(2, "8.2", "团队分工")
    b.body("项目组成员与职责分工如下表所示。")
    b.tcaption("表8-2 团队分工")
    b.table(["序号", "成员", "角色", "主要职责"], C.TABLE_8_2,
            [1.0, 1.8, 3.4, 10.79], size=8.5, align_center_cols=(0,))

    b.heading(2, "8.3", "需求变更管理")
    b.body_many(C.SEC8_CHANGE)


# ── 第 9 章 ─────────────────────────────────────────────────────────────────
def section9(b):
    b.heading(1, "9", "确认与签署", break_before=True)
    b.body_many(C.SEC9_CONFIRM)
    b.tcaption("表9-1 确认与签署")
    b.table(["角色", "姓名", "签字", "日期"], C.TABLE_9_1,
            [4.0, 3.0, 5.0, 4.99], align_center_cols=(0, 1))


# ── 入口：两遍生成，第二遍写入实测页码 ──────────────────────────────────────
def main():
    tmp = BASE / "_pass1.docx"
    toc = build(tmp)
    pages, total = BD.measure_pages(tmp, toc)
    print(f"pass1: {total} pages")
    build(OUT, toc_pages=pages, total_pages=total)

    # 校验第二遍页码是否漂移
    pages2, total2 = BD.measure_pages(OUT, toc)
    drift = {k: (pages[k], pages2[k]) for k in pages if pages[k] != pages2.get(k)}
    print(f"pass2: {total2} pages | 页码漂移: {drift if drift else '无'}")
    print("输出:", OUT)


if __name__ == "__main__":
    main()
